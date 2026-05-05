import pandas as pd
from io import BytesIO, StringIO
from ortools.sat.python import cp_model
from openpyxl import load_workbook

from .context import Context
from . import utils, models, constants


def get_people_versus_date_dataframe(ctx: Context, solver: cp_model.CpSolver, prettify: bool = False):
    # Initialize dataframe with size including leading rows and columns
    n_leading_rows, n_leading_cols = 2, 1
    n_trailing_rows, n_trailing_cols = 2, 0
    
    # Dictionary to track cells with [X] markers and their weights for Excel notes
    cell_export_info = {}
    
    n_history_cols = 0
    # Add history columns after the name column (only if prettify is enabled)
    if prettify:
        max_history_length = max((len(person.history) for person in ctx.people.items if person.history), default=0)
        n_history_cols = max_history_length
    
    # Add extra columns and rows for prettify mode
    shift_type_cols = ctx.n_shift_types + len(ctx.shiftTypes.groups) if prettify else 0  # Individual shift types + shift type groups
    extra_cols = (6 + shift_type_cols) if prettify else shift_type_cols  # Empty column + 5 OFF count columns + shift type columns
    extra_rows = (1 + ctx.n_shift_types + len(ctx.shiftTypes.groups)) if prettify else 0  # Empty row + one row per shift type + one row per shift type group
    
    df = pd.DataFrame(
        "",
        index=range(n_leading_rows + len(ctx.people.items) + n_trailing_rows + extra_rows),
        columns=range(n_leading_cols + n_history_cols + len(ctx.dates.items) + n_trailing_cols + extra_cols),
        dtype=object
    )

    # Fill history column headers (only if prettify is enabled)
    if n_history_cols > 0:
        # - row 0 contains history position labels (H-1, H-2, etc.)
        # - row 1 contains "History" label
        for h in range(n_history_cols):
            df.iloc[0, n_leading_cols + h] = f"H-{n_history_cols - h}"
            df.iloc[1, n_leading_cols + h] = "History"
    
    # Fill day numbers and weekdays
    # - row 0 contains day number
    # - row 1 contains weekday
    for d, date in enumerate(ctx.dates.items):
        col_idx = n_leading_cols + n_history_cols + d
        if ctx.dates.items[0].year != ctx.dates.items[-1].year:
            df.iloc[0, col_idx] = date.strftime('%Y/%-m/%-d')
        elif ctx.dates.items[0].month != ctx.dates.items[-1].month:
            df.iloc[0, col_idx] = date.strftime('%-m/%-d')
        else:
            df.iloc[0, col_idx] = date.day
        df.iloc[1, col_idx] = date.strftime('%a')

    # Fill person descriptions and history
    # - column 0 contains person description
    # - columns 1 to n_history_cols contain history data (padded with empty strings, only if prettify)
    for p, person in enumerate(ctx.people.items):
        df.iloc[n_leading_rows+p, 0] = person.id
        
        # Fill history columns with proper padding (only if prettify is enabled)
        if n_history_cols > 0:
            if person.history:
                history = person.history
                # Pad with empty strings at the front if history is shorter than max_history_length
                max_history_length = max((len(person.history) for person in ctx.people.items if person.history), default=0)
                padded_history = [""] * (max_history_length - len(history)) + history
                for h in range(n_history_cols):
                    df.iloc[n_leading_rows+p, n_leading_cols + h] = padded_history[h]
            else:
                # Fill with empty strings if no history
                for h in range(n_history_cols):
                    df.iloc[n_leading_rows+p, n_leading_cols + h] = ""

    # Pre-filter preferences to avoid repeated filtering in the inner loop
    filtered_preferences = {}
    if prettify:
        for pref in ctx.preferences:
            if pref.type != models.SHIFT_REQUEST:
                continue
            if pref.weight == 0:
                continue
            ds = utils.parse_dates(pref.date, ctx.map_did_d, ctx.dates.range)
            ss = utils.parse_sids(pref.shiftType, ctx.map_sid_s)
            ps = utils.parse_pids(pref.person, ctx.map_pid_p)
            if len(pref.shiftType) != 1 or len(ps) != 1:
                # Skip since is not single person and single shift type style
                continue
            if len(ds) != len(pref.date):
                # Skip since only count for single-date style
                continue
            
            # Store filtered preferences by (d, p) key for quick lookup
            for d in ds:
                for p in ps:
                    if (d, p) not in filtered_preferences:
                        filtered_preferences[(d, p)] = []
                    filtered_preferences[(d, p)].append({
                        'pref': pref,
                        'ss': ss,
                        'target_value': 1 if pref.weight > 0 else 0
                    })

    # Set cell values based on solver results
    for (d, p) in ctx.map_dp_s.keys():
        col_idx = n_leading_cols + n_history_cols + d
        assert df.iloc[n_leading_rows+p, col_idx] == ""
        cell_value = ""
        for s in ctx.map_dp_s[(d, p)]:
            if solver.Value(ctx.shifts[(d, s, p)]) == 1:
                if cell_value != "":
                    cell_value += ", "
                cell_value += ctx.shiftTypes.items[s].id
        if prettify:
            # Only consider single-person, single-shift-type, list-of-single-date style shift request
            # Add a ` [OFF]` suffix if the person requests OFF
            # Add a ` [<shift type id>]` suffix if the person requests a specific shift type
            # Add a ` [X]` suffix if the shift request is violated
            # Use pre-filtered preferences for this (d, p) combination
            if (d, p) in filtered_preferences:
                for pref_data in filtered_preferences[(d, p)]:
                    pref = pref_data['pref']
                    ss = pref_data['ss']
                    target_value = pref_data['target_value']
                    # Does not support shift type groups with mixed OFF and non-OFF shift types,
                    # which in most cases should not happen.
                    vars = [ctx.shifts[(d, s, p)] for s in ss] if constants.OFF_sid not in ss else [ctx.offs[(d, p)]]
                    if constants.OFF_sid in ss:
                        cell_value += " [OFF]"
                    else:
                        assert len(pref.shiftType) == 1
                        cell_value += f" [{pref.shiftType[0]}]"
                    if all((solver.Value(var) != target_value) for var in vars):
                        cell_value += " [X]"
                        # Track this cell for Excel notes - store the weight
                        excel_row = n_leading_rows + p + 1  # +1 for 1-based Excel indexing
                        excel_col = n_leading_cols + n_history_cols + d + 1  # +1 for 1-based Excel indexing
                        if (excel_row, excel_col) not in cell_export_info:
                            cell_export_info[(excel_row, excel_col)] = []
                        cell_export_info[(excel_row, excel_col)].append(abs(pref.weight))
        df.iloc[n_leading_rows+p, col_idx] = cell_value

    # Fill objective value
    df.iloc[n_leading_rows + len(ctx.people.items), 0] = "Score"
    df.iloc[n_leading_rows + len(ctx.people.items), n_leading_cols + n_history_cols] = solver.Value(ctx.objective) # or solver.ObjectiveValue()
    # Fill solver status
    df.iloc[n_leading_rows + len(ctx.people.items) + 1, 0] = "Status"
    df.iloc[n_leading_rows + len(ctx.people.items) + 1, n_leading_cols + n_history_cols] = ctx.solver_status

    # Sanity check with offs variables
    if not prettify:
        for (d, p) in ctx.offs.keys():
            col_idx = n_leading_cols + n_history_cols + d
            if solver.Value(ctx.offs[(d, p)]) == 1:
                assert df.iloc[n_leading_rows+p, col_idx] == ""
            else:
                assert df.iloc[n_leading_rows+p, col_idx] != ""

    if prettify:
        # TODO: Remove the concept of workday and freeday groups entirely from core solver
        # Instead, suggest these dates in the Web UI.
        # After that, we can get the ID with exact match.
        # Find date groups containing "workday" and "freeday" (case-insensitive)
        workday_group_id = None
        freeday_group_id = None
        
        # Use filter to simplify group lookup and ensure uniqueness
        workday_groups = list(filter(lambda g: "workday" in g.id.lower(), ctx.dates.groups))
        freeday_groups = list(filter(lambda g: "freeday" in g.id.lower(), ctx.dates.groups))
        workday_group_id = workday_groups[0].id if len(workday_groups) == 1 else None
        freeday_group_id = freeday_groups[0].id if len(freeday_groups) == 1 else None
        if not workday_group_id or not freeday_group_id:
            workday_group_id = None
            freeday_group_id = None
        
        # Add headers for OFF count columns
        off_col_start = n_leading_cols + n_history_cols + len(ctx.dates.items) + 1
        col_idx = off_col_start
        
        # Add workday and freeday columns if groups are found
        if workday_group_id and freeday_group_id:
            df.iloc[1, col_idx] = f"OFF ({workday_group_id})"
            col_idx += 1
            df.iloc[1, col_idx] = f"OFF ({freeday_group_id})"
            col_idx += 1
            
        df.iloc[1, col_idx] = "OFF (Total)"
        df.iloc[1, col_idx + 1] = "OFF (Weekday)"
        df.iloc[1, col_idx + 2] = "OFF (Weekend)"
        
        # Add headers for shift type count columns
        shift_type_col_start = col_idx + 3  # After the OFF columns
        col_idx = shift_type_col_start
        
        # Add individual shift type column headers
        for s in range(ctx.n_shift_types):
            df.iloc[1, col_idx] = f"{ctx.shiftTypes.items[s].id} Count"
            col_idx += 1
        
        # Add shift type group column headers
        for shift_group in ctx.shiftTypes.groups:
            df.iloc[1, col_idx] = f"{shift_group.id} Count"
            col_idx += 1
        
        # Count OFF days for each person (rows)
        for p in range(len(ctx.people.items)):
            off_workday = 0
            off_freeday = 0
            off_total = 0
            off_weekday = 0
            off_weekend = 0
            
            for d in range(len(ctx.dates.items)):
                if solver.Value(ctx.offs[(d, p)]) == 1:
                    off_total += 1
                    # Check if this date is a weekend
                    date = ctx.dates.items[d]
                    if date.weekday() >= 5:  # Saturday=5, Sunday=6
                        off_weekend += 1
                    else:
                        off_weekday += 1
                    
                    # Check if this date belongs to workday or freeday groups
                    if workday_group_id and freeday_group_id:
                        if d in ctx.map_did_d[workday_group_id]:
                            off_workday += 1
                        if d in ctx.map_did_d[freeday_group_id]:
                            off_freeday += 1
            
            # Fill the OFF count columns
            col_idx = off_col_start
            if workday_group_id and freeday_group_id:
                df.iloc[n_leading_rows + p, col_idx] = off_workday
                col_idx += 1
                df.iloc[n_leading_rows + p, col_idx] = off_freeday
                col_idx += 1
                
            df.iloc[n_leading_rows + p, col_idx] = off_total
            df.iloc[n_leading_rows + p, col_idx + 1] = off_weekday
            df.iloc[n_leading_rows + p, col_idx + 2] = off_weekend
            
            # Fill shift type count columns for this person
            shift_col_idx = shift_type_col_start
            
            # Count individual shift types for this person
            for s in range(ctx.n_shift_types):
                shift_count = sum(1 for d in range(len(ctx.dates.items))
                                if solver.Value(ctx.shifts[(d, s, p)]) == 1)
                df.iloc[n_leading_rows + p, shift_col_idx] = shift_count
                shift_col_idx += 1
            
            # Count shift type groups for this person
            for shift_group in ctx.shiftTypes.groups:
                shift_group_count = sum(1 for d in range(len(ctx.dates.items))
                                      if any(solver.Value(ctx.shifts[(d, s, p)]) == 1
                                           for member_id in shift_group.members
                                           for s in ctx.map_sid_s[member_id]))
                df.iloc[n_leading_rows + p, shift_col_idx] = shift_group_count
                shift_col_idx += 1
        
        # Add shift type count rows for each date (columns)
        # First add an empty row
        empty_row_index = n_leading_rows + len(ctx.people.items) + n_trailing_rows
        
        # Add one row for each individual shift type
        for s in range(ctx.n_shift_types):
            shift_row_index = empty_row_index + 1 + s
            df.iloc[shift_row_index, 0] = f"{ctx.shiftTypes.items[s].id} Count"
            
            for d in range(len(ctx.dates.items)):
                shift_count = sum(1 for p in range(len(ctx.people.items))
                                  if solver.Value(ctx.shifts[(d, s, p)]) == 1)
                df.iloc[shift_row_index, n_leading_cols + n_history_cols + d] = shift_count
        
        # Add one row for each shift type group
        for g, shift_group in enumerate(ctx.shiftTypes.groups):
            shift_row_index = empty_row_index + 1 + ctx.n_shift_types + g
            df.iloc[shift_row_index, 0] = f"{shift_group.id} Count"
            
            for d in range(len(ctx.dates.items)):
                shift_count = sum(1 for p in range(len(ctx.people.items))
                                  if any(solver.Value(ctx.shifts[(d, s, p)]) == 1
                                  for member_id in shift_group.members
                                  for s in ctx.map_sid_s[member_id]))
                df.iloc[shift_row_index, n_leading_cols + n_history_cols + d] = shift_count

    # Apply weekend highlighting and borders if prettify is enabled
    if prettify:
        # Create a styler object to apply conditional formatting
        def apply_styling(df):
            # Create a style DataFrame with the same shape as the original
            style_df = pd.DataFrame('', index=df.index, columns=df.columns)
            
            # Apply center alignment to all cells
            for row_idx in range(len(df)):
                for col_idx in range(len(df.columns)):
                    style_df.iloc[row_idx, col_idx] = 'text-align: center'
            
            # Apply dark red font color to cells containing violation markers "[X]"
            for row_idx in range(len(df)):
                for col_idx in range(len(df.columns)):
                    cell_value = df.iloc[row_idx, col_idx]
                    if cell_value and isinstance(cell_value, str) and "[X]" in cell_value:
                        existing_style = style_df.iloc[row_idx, col_idx]
                        if existing_style:
                            style_df.iloc[row_idx, col_idx] = f"{existing_style}; color: #C00000"
                        else:
                            style_df.iloc[row_idx, col_idx] = "color: #C00000"
            
            # Apply light yellow background to history columns
            for col_idx in range(n_leading_cols, n_leading_cols + n_history_cols):
                for row_idx in range(len(df)):
                    existing_style = style_df.iloc[row_idx, col_idx]
                    if existing_style:
                        style_df.iloc[row_idx, col_idx] = f"{existing_style}; background-color: #fefce8"
                    else:
                        style_df.iloc[row_idx, col_idx] = 'background-color: #fefce8'
            
            # Check each column to see if it represents a weekend or freeday group (only date columns, not history columns)
            for col_idx in range(n_leading_cols + n_history_cols, n_leading_cols + n_history_cols + len(ctx.dates.items)):
                d = col_idx - n_leading_cols - n_history_cols  # Get the date index
                
                # Get the weekday from row 1 (second row)
                weekday = df.iloc[1, col_idx]
                
                # Check if this date belongs to the freeday group (highest priority)
                if freeday_group_id and d in ctx.map_did_d[freeday_group_id]:
                    # Apply light green background for freeday group
                    for row_idx in range(len(df)):
                        existing_style = style_df.iloc[row_idx, col_idx]
                        if existing_style:
                            style_df.iloc[row_idx, col_idx] = f"{existing_style}; background-color: #dcfce7"
                        else:
                            style_df.iloc[row_idx, col_idx] = 'background-color: #dcfce7'
                # If it's Saturday or Sunday, highlight the entire column (lower priority than freeday)
                elif weekday in ['Sat', 'Sun']:
                    # Apply background color to all rows in this column
                    for row_idx in range(len(df)):
                        existing_style = style_df.iloc[row_idx, col_idx]
                        if existing_style:
                            style_df.iloc[row_idx, col_idx] = f"{existing_style}; background-color: #dbeafe"
                        else:
                            style_df.iloc[row_idx, col_idx] = 'background-color: #dbeafe'
            
            # Add borders to separate regions
            # Horizontal borders
            header_row_end = n_leading_rows - 1  # End of header region
            people_row_end = header_row_end + len(ctx.people.items)  # End of people region
            summary_row_end = people_row_end + n_trailing_rows  # End of summary region
            shift_type_individual_counts_row_end = summary_row_end + 1 + ctx.n_shift_types  # End of individual shift type counts
            shift_type_group_counts_row_end = shift_type_individual_counts_row_end + len(ctx.shiftTypes.groups)  # End of group counts
            
            # Vertical borders  
            name_col_end = n_leading_cols - 1  # End of name column
            history_col_end = name_col_end + n_history_cols  # End of history columns
            date_col_end = history_col_end + len(ctx.dates.items)  # End of date columns
            # Calculate OFF column positions dynamically
            off_total_col = date_col_end + 1  # Start after empty column
            if workday_group_id and freeday_group_id:
                off_total_col += 2
            off_total_col_end = off_total_col + 1  # End of OFF (Total) column
            # Calculate the number of remaining OFF columns
            num_remaining_off_cols = 2  # Weekday, Weekend
            off_weekend_col_end = off_total_col_end + num_remaining_off_cols  # End of last OFF count column
            # Calculate shift type column positions
            shift_type_individual_counts_col_end = off_weekend_col_end + ctx.n_shift_types
            shift_type_group_counts_col_end = shift_type_individual_counts_col_end + len(ctx.shiftTypes.groups)
            
            # Apply borders to all cells, then add specific border styles
            for row_idx in range(len(df)):
                for col_idx in range(len(df.columns)):
                    base_style = style_df.iloc[row_idx, col_idx]
                    borders = []
                    
                    # Add horizontal borders
                    if row_idx in [header_row_end, people_row_end, summary_row_end, shift_type_individual_counts_row_end, shift_type_group_counts_row_end]:
                        borders.append('border-bottom: 2px solid #374151')
                    
                    # Add vertical borders
                    if col_idx in [name_col_end, history_col_end, date_col_end, off_total_col_end, off_weekend_col_end, shift_type_individual_counts_col_end, shift_type_group_counts_col_end]:
                        borders.append('border-right: 2px solid #374151')
                    
                    # Add vertical border after Saturday columns (between Saturday and Sunday)
                    if (n_leading_cols + n_history_cols <= col_idx < n_leading_cols + n_history_cols + len(ctx.dates.items) and
                        df.iloc[1, col_idx] == 'Sat'):
                        borders.append('border-right: 2px solid #9ca3af')
                    
                    # Combine base style with borders
                    if borders:
                        border_style = '; '.join(borders)
                        if base_style:
                            style_df.iloc[row_idx, col_idx] = f"{base_style}; {border_style}"
                        else:
                            style_df.iloc[row_idx, col_idx] = border_style
            
            return style_df
        
        # Apply the styling and return the styled DataFrame
        styled_df = df.style.apply(lambda x: apply_styling(df), axis=None)
        return styled_df, cell_export_info
    
    return df, cell_export_info


def export_to_excel(df, output_buffer, cell_export_info=None):
    """
    Export DataFrame to Excel with frozen panes at B3 (first two rows and first column).
    Also adds notes/comments to cells with [X] markers showing the weight of unmet single-style requests.
    
    Args:
        output_buffer: BytesIO buffer to write to
        cell_export_info: Dictionary containing cell comment information
    """
    
    # Write to a temporary BytesIO buffer first
    temp_buffer = BytesIO()
    df.to_excel(temp_buffer, index=False, header=False)
    temp_buffer.seek(0)
    
    # Load the workbook to apply additional formatting
    wb = load_workbook(temp_buffer)
    ws = wb.active
    
    # Freeze the first two rows and first column (B3 is the cell after frozen area)
    ws.freeze_panes = 'B3'
    
    # Add notes/comments to cells with [X] markers if cell_export_info is provided
    if cell_export_info:
        from openpyxl.comments import Comment
        for (row, col), weights in cell_export_info.items():
            cell = ws.cell(row=row, column=col)
            # Calculate total weight and create note text
            total_weight = sum(weights)
            if len(weights) == 1:
                note_text = f"Weight of unmet single-style request: {total_weight}"
            else:
                note_text = f"Weights of unmet single-style requests: {total_weight} (individual weights: {', '.join(map(str, weights))})"
            
            # Create and add the comment
            comment = Comment(note_text, "Nurse Scheduling System")
            cell.comment = comment
    
    # Save to the output buffer
    wb.save(output_buffer)
    output_buffer.seek(0)


def export_to_csv(df, output_buffer):
    """
    Export DataFrame to CSV with UTF-8 BOM for Excel compatibility.
    
    Args:
        output_buffer: BytesIO buffer to write to (use BytesIO for proper encoding handling)
    """
    # Write CSV to a StringIO first to get text, then encode with BOM
    temp_buffer = StringIO()
    df.to_csv(temp_buffer, index=False, header=False)
    temp_buffer.seek(0)
    
    # Encode with UTF-8 BOM and write to output buffer
    csv_content = temp_buffer.getvalue()
    output_buffer.write(csv_content.encode('utf-8-sig'))
    output_buffer.seek(0)
